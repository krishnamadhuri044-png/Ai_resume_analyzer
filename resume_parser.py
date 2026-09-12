"""
resume_parser.py
-----------------
Extracts raw text from uploaded PDF or DOCX resumes.
"""

import os
from pypdf import PdfReader
import docx


class UnsupportedFileTypeError(Exception):
    pass


def extract_text_from_pdf(file_path_or_buffer) -> str:
    """Extract text from every page of a PDF file."""
    reader = PdfReader(file_path_or_buffer)
    text_chunks = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        text_chunks.append(page_text)
    return "\n".join(text_chunks)


def extract_text_from_docx(file_path_or_buffer) -> str:
    """Extract text from every paragraph (and table cell) of a DOCX file."""
    document = docx.Document(file_path_or_buffer)
    text_chunks = [p.text for p in document.paragraphs if p.text]

    # Also capture text inside tables (many resumes use table layouts)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    text_chunks.append(cell.text)

    return "\n".join(text_chunks)


def extract_resume_text(file_path_or_buffer, filename: str = None) -> str:
    """
    Detect file type from extension and extract text accordingly.

    Parameters
    ----------
    file_path_or_buffer : str path or file-like object (e.g. Streamlit UploadedFile)
    filename : optional filename to use for extension detection when a
               buffer (not a path) is passed in.
    """
    name = filename or (file_path_or_buffer if isinstance(file_path_or_buffer, str) else "")
    ext = os.path.splitext(name)[1].lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_path_or_buffer)
    elif ext == ".docx":
        return extract_text_from_docx(file_path_or_buffer)
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext}'. Please upload a PDF or DOCX resume."
        )


def validate_file(filename: str, file_size_bytes: int, max_size_mb: int = 5) -> tuple:
    """
    Validate file type and size before parsing.

    Returns (is_valid: bool, message: str)
    """
    ext = os.path.splitext(filename)[1].lower()
    if ext not in (".pdf", ".docx"):
        return False, "Only PDF and DOCX files are supported."

    max_bytes = max_size_mb * 1024 * 1024
    if file_size_bytes > max_bytes:
        return False, f"File exceeds the {max_size_mb}MB size limit."

    return True, "OK"


if __name__ == "__main__":
    # Simple manual test
    import sys
    if len(sys.argv) > 1:
        path = sys.argv[1]
        print(extract_resume_text(path, filename=path)[:500])
    else:
        print("Usage: python resume_parser.py <path_to_resume.pdf|.docx>")
